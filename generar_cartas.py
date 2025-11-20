# generar_cartas.py
# Utilidades para generación de cartas y sílabos desde la BD

from datetime import date, datetime
import re
import unicodedata
from typing import Optional, Union, Tuple
import os

from docx import Document  # para los sílabos (python-docx)
from config import PLANTILLAS_SILABO_DIR, CARPETA_SILABOS

# ============================
# 1) Fecha larga en español
# ============================

MESES_ES = [
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "setiembre",
    "octubre",
    "noviembre",
    "diciembre",
]


def fecha_larga_es(
    fecha: Optional[Union[date, datetime]] = None,
) -> Tuple[str, str]:
    """
    Devuelve una fecha larga en español y el año, por ejemplo:
        ("14 de noviembre de 2025", "2025")

    - Si no se pasa fecha, usa la fecha actual.
    - Si se pasa datetime, usa solo la parte de fecha.
    """
    if fecha is None:
        hoy = date.today()
    elif isinstance(fecha, datetime):
        hoy = fecha.date()
    elif isinstance(fecha, date):
        hoy = fecha
    else:
        hoy = date.today()

    mes_nombre = MESES_ES[hoy.month - 1]
    fecha_larga = f"{hoy.day} de {mes_nombre} de {hoy.year}"
    return fecha_larga, str(hoy.year)


# ============================
# 2) Remuneración
# ============================

def _parse_monto(monto_str: Optional[str]) -> Optional[float]:
    """
    Convierte strings como:
      '5900'
      '5900.0'
      'S/. 5,900.00'
      's/ 5900'
    a float 5900.0

    Devuelve None si no se puede interpretar.
    """
    if not monto_str:
        return None

    s = str(monto_str).strip()

    if s in ("-", "--", ""):
        return None

    # Quitar prefijos de moneda y espacios
    for pref in ["S/.", "S/", "s/.", "s/"]:
        s = s.replace(pref, "")
    s = s.replace(" ", "")

    # Quitar separadores de miles
    s = s.replace(",", "")

    try:
        return float(s)
    except ValueError:
        return None


def calc_remuneracion(
    programa_texto: str,
    override: Optional[str] = None,
) -> str:
    """
    Lógica:

    - Si 'override' viene con valor (desde la BD), se usa ese monto.
    - Si NO hay override:
        * Si el texto del programa contiene 'doctorado' => 5900.00
        * Si contiene 'maestría' o 'maestria' => 4140.00
        * En otro caso => 0.00

    Retorna SIEMPRE un string:
      '5,900.00', '4,140.00', '0.00'
    """
    monto = _parse_monto(override)
    if monto is not None:
        return f"{monto:,.2f}"

    t = (programa_texto or "").lower()

    if "doctorado" in t:
        monto = 5900.00
    elif "maestría" in t or "maestria" in t:
        monto = 4140.00
    else:
        monto = 0.0

    return f"{monto:,.2f}"


# ============================
# 3) Normalizar nombre de archivo
# ============================

def _strip_accents(text: str) -> str:
    """
    Elimina acentos y caracteres combinados.
    """
    text_norm = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text_norm if not unicodedata.combining(ch))


def limpiar_nombre_archivo(nombre: str) -> str:
    """
    Limpia un texto para usarlo como nombre de archivo.
    """
    if not nombre:
        return "archivo"

    nombre = nombre.strip()
    nombre = _strip_accents(nombre)

    # Dejar solo letras, números, espacios, guiones y puntos
    nombre = re.sub(r"[^A-Za-z0-9\.\- ]+", "", nombre)

    # Reemplazar espacios por guiones bajos
    nombre = nombre.replace(" ", "_")

    # Compactar guiones bajos repetidos
    nombre = re.sub(r"_+", "_", nombre)

    # Quitar guiones bajos iniciales/finales
    nombre = nombre.strip("_")

    return nombre or "archivo"


# ============================
# 4) Títulos y nombres
# ============================

_TITULOS_INICIALES = {
    "DR", "DRA",
    "MG", "MGS", "MBA", "MSC",
    "LIC", "ING", "ARQ",
    "MEDICO", "MÉDICO", "PSIC",
    "ABOG", "ABOGADO", "ABOGADA",
}


def separar_titulo_y_nombre(nombre_completo: str) -> Tuple[str, str]:
    """
    Separa el título (Dra., Dr., Mg., Lic., etc.) del resto del nombre.
    """
    if not nombre_completo:
        return "", "DOCENTE"

    partes = nombre_completo.strip().split()
    if not partes:
        return "", "DOCENTE"

    titulo = ""
    token_raw = partes[0]
    token = token_raw.upper().rstrip(".")

    if token in _TITULOS_INICIALES:
        titulo = token_raw.rstrip(".") + "."
        partes = partes[1:]

    nombre_sin_titulo = " ".join(partes) if partes else "DOCENTE"
    return titulo, nombre_sin_titulo


def extraer_paterno_y_nombre(nombre_completo: str) -> Tuple[str, str]:
    """
    Devuelve (paterno, nombre) sin títulos.
    """
    if not nombre_completo:
        return "DOCENTE", ""

    partes = nombre_completo.strip().split()

    while partes:
        token = partes[0].upper().rstrip(".")
        if token in _TITULOS_INICIALES:
            partes.pop(0)
        else:
            break

    if not partes:
        return "DOCENTE", ""

    if len(partes) == 1:
        return partes[0].upper(), ""
    if len(partes) == 2:
        return partes[0].upper(), partes[1].upper()

    paterno = partes[0].upper()
    nombre = partes[-1].upper()
    return paterno, nombre


def extraer_titulo_paterno_nombre(nombre_completo: str) -> Tuple[str, str, str]:
    """
    'DRA. MIRIAM CLEDY ZARATE MUÑIZ' -> ('DRA', 'ZARATE', 'MIRIAM')
    """
    if not nombre_completo:
        return "", "DOCENTE", ""

    partes = nombre_completo.strip().split()

    titulos = []
    while partes:
        token_raw = partes[0]
        token_norm = token_raw.upper().rstrip(".")
        if token_norm in _TITULOS_INICIALES:
            titulos.append(token_norm)
            partes.pop(0)
        else:
            break

    titulo = " ".join(titulos)

    if not partes:
        if not titulo:
            return "", "DOCENTE", ""
        return titulo, "DOCENTE", ""

    if len(partes) == 1:
        paterno = partes[0].upper()
        nombre = ""
    else:
        paterno = partes[-2].upper()
        nombre = partes[0].upper()

    return titulo, paterno, nombre


# =====================================================
# 5) UTILIDADES PARA GENERAR SÍLABOS
# =====================================================

def _reemplazar_tokens_en_elemento(elemento, mapping: dict) -> None:
    """
    Reemplaza tokens {{TOKEN}} en párrafos y tablas
    de un elemento de python-docx (Document o Cell).
    """
    # Párrafos
    for p in elemento.paragraphs:
        if not p.text:
            continue

        texto_original = p.text
        texto_nuevo = texto_original
        tiene_cambios = False

        for clave, valor in mapping.items():
            if clave in texto_nuevo:
                texto_nuevo = texto_nuevo.replace(clave, valor)
                tiene_cambios = True

        if tiene_cambios:
            if p.runs:
                p.runs[0].text = texto_nuevo
                for run in p.runs[1:]:
                    run.text = ""
            else:
                p.text = texto_nuevo

    # Tablas (recursivo en celdas)
    for tabla in elemento.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                _reemplazar_tokens_en_elemento(celda, mapping)


def generar_silabo_curso(db, curso_row):
    """
    Genera el sílabo (.docx) para un curso_programado.

    Parámetros:
        db         -> conexión (get_db_connection())
        curso_row  -> fila de SQLite con, al menos:
                      - id
                      - programa_id
                      - asignatura
                      - codigo
                      - categoria
                      - ciclo
                      - modalidad_dictado
                      - periodo_academico
                      - programa_nombre
                      - docente_nombre
                      - docente_correo

    Retorna:
        ruta completa del archivo .docx generado.
    """
    # 1) Configuración del programa
    prog_conf = db.execute(
        """
        SELECT *
        FROM silabo_programa_config
        WHERE programa_id = ?
        """,
        (curso_row["programa_id"],),
    ).fetchone()

    if not prog_conf:
        raise RuntimeError(
            "No hay configuración de sílabo para este programa (silabo_programa_config)."
        )

    # 2) Sumilla base (por código o por asignatura)
    sumilla = ""

    fila_sumilla = None
    if curso_row["codigo"]:
        fila_sumilla = db.execute(
            """
            SELECT sumilla
            FROM silabo_curso_base
            WHERE programa_id = ?
              AND codigo = ?
            """,
            (curso_row["programa_id"], curso_row["codigo"]),
        ).fetchone()

    if not fila_sumilla and curso_row["asignatura"]:
        fila_sumilla = db.execute(
            """
            SELECT sumilla
            FROM silabo_curso_base
            WHERE programa_id = ?
              AND UPPER(asignatura) = UPPER(?)
            """,
            (curso_row["programa_id"], curso_row["asignatura"]),
        ).fetchone()

    if fila_sumilla:
        sumilla = fila_sumilla["sumilla"]

    # 3) Derivar textos
    periodo_acad = curso_row["periodo_academico"] or ""
    modalidad = (curso_row["modalidad_dictado"] or prog_conf["modalidad_default"] or "").upper()

    horas_tp = f"{prog_conf['horas_teoricas']} HT – {prog_conf['horas_practicas']} HP"
    num_cred = str(prog_conf["numero_creditos"])

    if prog_conf["inicio_semestre"] and prog_conf["fin_semestre"]:
        inicio_fin = f"{prog_conf['inicio_semestre']} al {prog_conf['fin_semestre']}"
    else:
        anio = periodo_acad.split("-")[0] if "-" in periodo_acad else ""
        inicio_fin = anio

    horario = prog_conf["horario_texto"] or (curso_row["horas_texto"] or "")

    docente_nombre = curso_row["docente_nombre"] or ""
    docente_correo = curso_row["docente_correo"] or ""

    asignatura = curso_row["asignatura"] or ""
    asignatura_titulo = asignatura.upper()

    # 4) Determinar plantilla
    plantilla_archivo = prog_conf["plantilla_archivo"] or "plantilla_silabo_generica.docx"
    plantilla_path = PLANTILLAS_SILABO_DIR / plantilla_archivo

    if not plantilla_path.exists():
        raise RuntimeError(f"No se encontró la plantilla de sílabo: {plantilla_path}")

    doc = Document(str(plantilla_path))

    # 5) Mapeo de tokens -> valores
    mapping = {
        "{{NOMBRE_PROGRAMA}}": prog_conf["nombre_programa"],
        "{{MODALIDAD_TITULO}}": modalidad,
        "{{ASIGNATURA_TITULO}}": asignatura_titulo,

        "{{ASIGNATURA}}": asignatura,
        "{{CODIGO_ASIGNATURA}}": curso_row["codigo"] or "",
        "{{CATEGORIA}}": curso_row["categoria"] or "",
        "{{SEMESTRE_ACADEMICO}}": periodo_acad,
        "{{CICLO}}": curso_row["ciclo"] or "",
        "{{NUM_CREDITOS}}": num_cred,
        "{{HORAS_TP}}": horas_tp,
        "{{MODALIDAD}}": modalidad,
        "{{HORARIO}}": horario,
        "{{INICIO_FIN_SEMESTRE}}": inicio_fin,
        "{{DOCENTE_NOMBRE}}": docente_nombre,
        "{{DOCENTE_CORREO}}": docente_correo,
        "{{MEET_LINK}}": "",  # si luego guardas enlace meet en BD, se rellena aquí

        "{{SUMILLA}}": sumilla,
        "{{PERFIL_EGRESADO}}": prog_conf["perfil_egresado"] or "",
        "{{RESULTADOS_APRENDIZAJE}}": prog_conf["resultados_aprendizaje"] or "",
    }

    # 6) Reemplazar tokens
    _reemplazar_tokens_en_elemento(doc, mapping)

    # 7) Nombre de archivo
    codigo = curso_row["codigo"] or "SIN_CODIGO"
    docente_slug = limpiar_nombre_archivo(docente_nombre) if docente_nombre else "SIN_DOCENTE"

    nombre_archivo = f"SILABO_{codigo}_{docente_slug}.docx"
    output_path = CARPETA_SILABOS / nombre_archivo

    doc.save(str(output_path))
    return str(output_path)
