from flask import (
    Flask,
    render_template,
    request,
    Response,
    redirect,
    url_for,
    flash,
    send_file,
)
import sqlite3
from pathlib import Path
import os
from dotenv import load_dotenv
import csv
from io import StringIO, BytesIO
from docx import Document  # para exportar a Word (listado docentes)
from docx.shared import Cm
from docx.enum.section import WD_ORIENT
import re
from reporte_docentes_sunedu import register_sunedu_routes

# === IMPORTS NUEVOS PARA CARTAS DE INVITACIÓN ===
from docxtpl import DocxTemplate
from generar_cartas import (
    fecha_larga_es,
    calc_remuneracion,
    limpiar_nombre_archivo,
    extraer_paterno_y_nombre,
)
# PDF opcional (como en generar_cartas.py)
TRY_PDF_WEB = True
try:
    from docx2pdf import convert as docx2pdf_convert
except Exception:
    TRY_PDF_WEB = False
    
# === Configuración base ===
BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / ".env"
if ENV_PATH.exists():
    load_dotenv(ENV_PATH)

DB_PATH = os.getenv("DOCENTES_DB_PATH", str(BASE_DIR / "docentes.db"))

# Carpeta base donde se guardarán las cartas generadas desde la web
CARTAS_OUTPUT_DIR = BASE_DIR / "cartas_invitacion"
CARTAS_OUTPUT_DIR.mkdir(exist_ok=True)

# Plantilla de la carta de invitación
PLANTILLA_CARTA = BASE_DIR / "plantilla_carta.docx"

app = Flask(__name__)
app.secret_key = "cambia-esto-por-algo-mas-seguro"  # para mensajes flash


def get_db_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

# ⬇️ NUEVO: registrar rutas del reporte SUNEDU
register_sunedu_routes(app, get_db_connection)

# ====== Datos comunes inyectados a todas las vistas ======
@app.context_processor
def inject_common_filters():
    conn = get_db_connection()
    facultades = conn.execute(
        "SELECT id, nombre FROM facultad WHERE activo = 1 ORDER BY nombre"
    ).fetchall()

    periodos = conn.execute(
        "SELECT id, etiqueta FROM periodo ORDER BY anio DESC, mes DESC"
    ).fetchall()
    conn.close()

    tipos_programa = [
        ("", "Todos"),
        ("DOCTORADO", "Doctorado"),
        ("MAESTRIA", "Maestría"),
    ]

    tipos_docente_mes = [
        ("", "Todos"),
        ("LOCAL", "Local"),
        ("ORDINARIZADO", "Ordinarizado"),
    ]

    return dict(
        facultades_global=facultades,
        periodos_global=periodos,
        tipos_programa_global=tipos_programa,
        tipos_docente_mes_global=tipos_docente_mes,
    )


# ====== Dashboard ======
@app.route("/")
def index():
    conn = get_db_connection()

    periodo = conn.execute(
        "SELECT id, etiqueta, anio, mes FROM periodo ORDER BY anio DESC, mes DESC LIMIT 1"
    ).fetchone()

    if periodo:
        periodo_id = periodo["id"]
        etiqueta_periodo = periodo["etiqueta"]
    else:
        periodo_id = None
        etiqueta_periodo = "Sin periodos"

    total_docentes = conn.execute(
        "SELECT COUNT(*) as c FROM docente WHERE activo = 1"
    ).fetchone()["c"]

    if periodo_id is not None:
        total_cursos = conn.execute(
            "SELECT COUNT(*) as c FROM curso_programado WHERE periodo_id = ?",
            (periodo_id,),
        ).fetchone()["c"]

        total_rem = conn.execute(
            "SELECT COALESCE(SUM(remuneracion_monto), 0) as total "
            "FROM curso_programado WHERE periodo_id = ?",
            (periodo_id,),
        ).fetchone()["total"]
    else:
        total_cursos = 0
        total_rem = 0

    conn.close()

    return render_template(
        "index.html",
        periodo=periodo,
        etiqueta_periodo=etiqueta_periodo,
        total_docentes=total_docentes,
        total_cursos=total_cursos,
        total_rem=total_rem,
    )


# ===================================================================
# 1) PERÍODOS – listar, crear, editar, eliminar
# ===================================================================

@app.route("/periodos")
def periodos_list():
    conn = get_db_connection()
    periodos = conn.execute(
        """
        SELECT p.*,
               (
                 SELECT COUNT(*) FROM curso_programado cp
                 WHERE cp.periodo_id = p.id
               ) AS total_cursos
        FROM periodo p
        ORDER BY p.anio DESC, p.mes DESC
        """
    ).fetchall()
    conn.close()
    return render_template("periodos_list.html", periodos=periodos)


@app.route("/periodos/nuevo", methods=["GET", "POST"])
def periodo_nuevo():
    if request.method == "POST":
        anio = request.form.get("anio", "").strip()
        mes = request.form.get("mes", "").strip()
        etiqueta = request.form.get("etiqueta", "").strip()

        if not anio or not mes or not etiqueta:
            flash("Año, mes y etiqueta son obligatorios", "danger")
            return render_template("periodo_form.html", periodo=None)

        try:
            anio_int = int(anio)
            mes_int = int(mes)
            if mes_int < 1 or mes_int > 12:
                raise ValueError
        except ValueError:
            flash("Mes debe ser un número entre 1 y 12", "danger")
            return render_template("periodo_form.html", periodo=None)

        conn = get_db_connection()
        try:
            conn.execute(
                "INSERT INTO periodo (anio, mes, etiqueta) VALUES (?, ?, ?)",
                (anio_int, mes_int, etiqueta),
            )
            conn.commit()
            flash("Período creado correctamente ✅", "success")
            return redirect(url_for("periodos_list"))
        except sqlite3.IntegrityError:
            flash("Ya existe un período con ese año y mes", "danger")
            return render_template("periodo_form.html", periodo=None)
        finally:
            conn.close()

    return render_template("periodo_form.html", periodo=None)


@app.route("/periodos/<int:periodo_id>/editar", methods=["GET", "POST"])
def periodo_editar(periodo_id):
    conn = get_db_connection()
    periodo = conn.execute(
        "SELECT * FROM periodo WHERE id = ?", (periodo_id,)
    ).fetchone()
    if not periodo:
        conn.close()
        flash("Período no encontrado", "danger")
        return redirect(url_for("periodos_list"))

    if request.method == "POST":
        anio = request.form.get("anio", "").strip()
        mes = request.form.get("mes", "").strip()
        etiqueta = request.form.get("etiqueta", "").strip()

        if not anio or not mes or not etiqueta:
            flash("Año, mes y etiqueta son obligatorios", "danger")
            return render_template("periodo_form.html", periodo=periodo)

        try:
            anio_int = int(anio)
            mes_int = int(mes)
            if mes_int < 1 or mes_int > 12:
                raise ValueError
        except ValueError:
            flash("Mes debe ser un número entre 1 y 12", "danger")
            return render_template("periodo_form.html", periodo=periodo)

        try:
            conn.execute(
                "UPDATE periodo SET anio = ?, mes = ?, etiqueta = ? WHERE id = ?",
                (anio_int, mes_int, etiqueta, periodo_id),
            )
            conn.commit()
            flash("Período actualizado correctamente ✅", "success")
            return redirect(url_for("periodos_list"))
        except sqlite3.IntegrityError:
            flash("Ya existe un período con ese año y mes", "danger")
            return render_template("periodo_form.html", periodo=periodo)
        finally:
            conn.close()

    conn.close()
    return render_template("periodo_form.html", periodo=periodo)


@app.route("/periodos/<int:periodo_id>/eliminar", methods=["POST"])
def periodo_eliminar(periodo_id):
    conn = get_db_connection()
    cursos = conn.execute(
        "SELECT COUNT(*) as c FROM curso_programado WHERE periodo_id = ?",
        (periodo_id,),
    ).fetchone()["c"]

    if cursos > 0:
        conn.close()
        flash(
            "No se puede eliminar el período porque tiene cursos programados.",
            "danger",
        )
        return redirect(url_for("periodos_list"))

    conn.execute("DELETE FROM periodo WHERE id = ?", (periodo_id,))
    conn.commit()
    conn.close()
    flash("Período eliminado correctamente ✅", "success")
    return redirect(url_for("periodos_list"))


# ===================================================================
# 2) DOCENTES – listado + CRUD básico + propuestos
# ===================================================================

@app.route("/docentes")
def docentes_list():
    tipo_docente = request.args.get("tipo_docente", "").strip()
    q = request.args.get("q", "").strip()
    periodo_propuesto_id = request.args.get("periodo_propuesto_id", "").strip()
    solo_con_cv = request.args.get("solo_con_cv", "") == "1"
    solo_propuestos = request.args.get("solo_propuestos", "") == "1"

    conn = get_db_connection()

    periodo_actual = conn.execute(
        "SELECT id FROM periodo ORDER BY anio DESC, mes DESC LIMIT 1"
    ).fetchone()
    periodo_actual_id = periodo_actual["id"] if periodo_actual else None

    periodo_cursos_filter = ""
    params = []

    if periodo_actual_id is not None:
        periodo_cursos_filter = "AND cp.periodo_id = ?"
        params.append(periodo_actual_id)

    if periodo_propuesto_id:
        sub_existe_propuesto = """
          CASE WHEN EXISTS(
            SELECT 1 FROM docente_propuesto dp
            WHERE dp.docente_id = d.id
              AND dp.periodo_id = ?
          ) THEN 1 ELSE 0 END AS es_propuesto
        """
        params.insert(0, periodo_propuesto_id)
        existe_condicion = """
          EXISTS(
            SELECT 1 FROM docente_propuesto dp2
            WHERE dp2.docente_id = d.id
              AND dp2.periodo_id = ?
          )
        """
    else:
        sub_existe_propuesto = """
          CASE WHEN EXISTS(
            SELECT 1 FROM docente_propuesto dp
            WHERE dp.docente_id = d.id
          ) THEN 1 ELSE 0 END AS es_propuesto
        """
        existe_condicion = """
          EXISTS(
            SELECT 1 FROM docente_propuesto dp2
            WHERE dp2.docente_id = d.id
          )
        """

    base_query = f"""
        SELECT d.*,
               (
                 SELECT COUNT(*)
                 FROM curso_programado cp
                 WHERE cp.docente_id = d.id
                 {periodo_cursos_filter}
               ) AS cantidad_cursos,
               {sub_existe_propuesto}
        FROM docente d
        WHERE d.activo = 1
    """

    if tipo_docente:
        base_query += " AND d.tipo_docente = ?"
        params.append(tipo_docente)

    if q:
        base_query += " AND (d.nombre_completo LIKE ? OR d.dni LIKE ?)"
        like_q = f"%{q}%"
        params.extend([like_q, like_q])

    if solo_con_cv:
        base_query += " AND d.tiene_cv = 1"

    if solo_propuestos:
        base_query += f" AND {existe_condicion}"
        if periodo_propuesto_id:
            params.append(periodo_propuesto_id)

    base_query += " ORDER BY d.nombre_completo ASC"

    docentes = conn.execute(base_query, params).fetchall()
    conn.close()

    return render_template(
        "docentes_list.html",
        docentes=docentes,
        tipo_docente=tipo_docente,
        q=q,
        periodo_propuesto_id=periodo_propuesto_id,
        solo_con_cv=solo_con_cv,
        solo_propuestos=solo_propuestos,
    )


@app.route("/docentes/nuevo", methods=["GET", "POST"])
def docente_nuevo():
    if request.method == "POST":
        nombre_completo = request.form.get("nombre_completo", "").strip()
        dni = request.form.get("dni", "").strip() or None
        especialidad = request.form.get("especialidad", "").strip()
        telefono = request.form.get("telefono", "").strip()
        correo = request.form.get("correo", "").strip()
        tipo_docente = request.form.get("tipo_docente", "").strip() or None
        tiene_cv = 1 if request.form.get("tiene_cv") == "1" else 0
        link_cv = request.form.get("link_cv", "").strip() or None
        fecha_recepcion_cv = request.form.get("fecha_recepcion_cv", "").strip() or None

        # Ahora SOLO nombre completo es obligatorio
        if not nombre_completo:
            flash("Nombre completo es obligatorio", "danger")
            return render_template("docente_form.html", docente=None)

        conn = get_db_connection()
        try:
            conn.execute(
                """
                INSERT INTO docente (
                    nombre_completo, dni, especialidad, telefono, correo,
                    tipo_docente, tiene_cv, link_cv, fecha_recepcion_cv
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    nombre_completo,
                    dni,
                    especialidad,
                    telefono,
                    correo,
                    tipo_docente,
                    tiene_cv,
                    link_cv,
                    fecha_recepcion_cv,
                ),
            )
            conn.commit()
            flash("Docente creado correctamente ✅", "success")
            return redirect(url_for("docentes_list"))
        except sqlite3.IntegrityError:
            # Esto solo saltaría si el DNI no es NULL y ya existe en otro docente
            flash("Ya existe un docente con ese DNI", "danger")
            return render_template("docente_form.html", docente=None)
        finally:
            conn.close()

    return render_template("docente_form.html", docente=None)


@app.route("/docentes/<int:docente_id>/editar", methods=["GET", "POST"])
def docente_editar(docente_id):
    conn = get_db_connection()
    docente = conn.execute(
        "SELECT * FROM docente WHERE id = ?", (docente_id,)
    ).fetchone()
    if not docente:
        conn.close()
        flash("Docente no encontrado", "danger")
        return redirect(url_for("docentes_list"))

    if request.method == "POST":
        nombre_completo = request.form.get("nombre_completo", "").strip()
        dni = request.form.get("dni", "").strip() or None
        especialidad = request.form.get("especialidad", "").strip()
        telefono = request.form.get("telefono", "").strip()
        correo = request.form.get("correo", "").strip()
        tipo_docente = request.form.get("tipo_docente", "").strip() or None
        tiene_cv = 1 if request.form.get("tiene_cv") == "1" else 0
        link_cv = request.form.get("link_cv", "").strip() or None
        fecha_recepcion_cv = request.form.get("fecha_recepcion_cv", "").strip() or None

        if not nombre_completo:
            flash("Nombre completo es obligatorio", "danger")
            return render_template("docente_form.html", docente=docente)

        try:
            conn.execute(
                """
                UPDATE docente
                SET nombre_completo = ?,
                    dni = ?,
                    especialidad = ?,
                    telefono = ?,
                    correo = ?,
                    tipo_docente = ?,
                    tiene_cv = ?,
                    link_cv = ?,
                    fecha_recepcion_cv = ?,
                    updated_at = datetime('now')
                WHERE id = ?
                """,
                (
                    nombre_completo,
                    dni,
                    especialidad,
                    telefono,
                    correo,
                    tipo_docente,
                    tiene_cv,
                    link_cv,
                    fecha_recepcion_cv,
                    docente_id,
                ),
            )
            conn.commit()
            flash("Docente actualizado correctamente ✅", "success")
            return redirect(url_for("docentes_list"))
        except sqlite3.IntegrityError:
            flash("Ya existe un docente con ese DNI", "danger")
            return render_template("docente_form.html", docente=docente)
        finally:
            conn.close()

    conn.close()
    return render_template("docente_form.html", docente=docente)


@app.route("/docentes/<int:docente_id>/toggle_propuesto", methods=["POST"])
def toggle_docente_propuesto(docente_id):
    periodo_id = request.form.get("periodo_id")
    if not periodo_id:
        flash("Periodo no especificado para propuesto", "danger")
        return redirect(url_for("docentes_list"))

    q = request.form.get("q", "")
    tipo_docente = request.form.get("tipo_docente", "")
    periodo_propuesto_id = request.form.get("periodo_propuesto_id", "")
    solo_con_cv = request.form.get("solo_con_cv", "")
    solo_propuestos = request.form.get("solo_propuestos", "")

    conn = get_db_connection()
    existe = conn.execute(
        "SELECT id FROM docente_propuesto WHERE docente_id = ? AND periodo_id = ?",
        (docente_id, periodo_id),
    ).fetchone()

    if existe:
        conn.execute(
            "DELETE FROM docente_propuesto WHERE docente_id = ? AND periodo_id = ?",
            (docente_id, periodo_id),
        )
        msg = "Docente quitado de propuestos para ese período."
    else:
        conn.execute(
            "INSERT INTO docente_propuesto (docente_id, periodo_id, estado) VALUES (?, ?, 'PROPUESTO')",
            (docente_id, periodo_id),
        )
        msg = "Docente marcado como propuesto para ese período."
    conn.commit()
    conn.close()

    flash(msg, "info")

    return redirect(
        url_for(
            "docentes_list",
            q=q,
            tipo_docente=tipo_docente,
            periodo_propuesto_id=periodo_propuesto_id,
            solo_con_cv=solo_con_cv,
            solo_propuestos=solo_propuestos,
        )
    )


# ===================================================================
# 3) PROGRAMACIÓN – listado, crear, editar, eliminar
# ===================================================================

@app.route("/programacion", methods=["GET"])
def programacion():
    conn = get_db_connection()
    cur = conn.cursor()

    # =========================
    # Catálogos para filtros
    # =========================
    cur.execute("SELECT id, etiqueta FROM periodo ORDER BY anio DESC, mes DESC")
    periodos_global = cur.fetchall()

    cur.execute("SELECT id, nombre FROM facultad WHERE activo = 1 ORDER BY nombre")
    facultades_global = cur.fetchall()

    tipos_programa_global = [
        ("", "Todos"),
        ("DOCTORADO", "Doctorados"),
        ("MAESTRIA", "Maestrías"),
    ]

    tipos_docente_mes_global = [
        ("", "Todos"),
        ("LOCAL", "Local"),
        ("ORDINARIZADO", "Ordinarizado"),
    ]

    # =========================
    # Lectura de filtros
    # =========================
    periodo_id = request.args.get("periodo_id", type=int)
    facultad_id = request.args.get("facultad_id", type=int)
    tipo_programa = request.args.get("tipo_programa", default="", type=str)
    tipo_docente_mes = request.args.get("tipo_docente_mes", default="", type=str)

    # Si no mandan período, tomamos el último disponible
    if not periodo_id and periodos_global:
        periodo_id = periodos_global[0]["id"]

    periodo_info = None
    if periodo_id:
        cur.execute("SELECT * FROM periodo WHERE id = ?", (periodo_id,))
        periodo_info = cur.fetchone()

    # =========================
    # Construcción dinámica del WHERE
    # =========================
    filtros = []
    params = []

    if periodo_id:
        filtros.append("cp.periodo_id = ?")
        params.append(periodo_id)

    if facultad_id:
        filtros.append("pa.facultad_id = ?")
        params.append(facultad_id)

    if tipo_programa:
        filtros.append("pa.tipo = ?")
        params.append(tipo_programa)

    if tipo_docente_mes:
        filtros.append("cp.tipo_docente_mes = ?")
        params.append(tipo_docente_mes)

    where_clause = " AND ".join(filtros) if filtros else "1 = 1"

    # =========================
    # Consulta principal de cursos
    # =========================
    sql_cursos = f"""
        SELECT
          cp.id AS curso_id,
          f.nombre AS facultad_nombre,
          pa.tipo AS programa_tipo,
          pa.nombre_corto AS programa_nombre,
          COALESCE(d.nombre_completo, '') AS docente_nombre,
          cp.dni_docente,
          cp.tipo_docente_mes,
          cp.ciclo,
          cp.asignatura,
          cp.fechas_texto,
          cp.remuneracion_texto,
          cp.poi,
          cp.observaciones,             -- 🔴 Observaciones en el listado
          pa.id AS programa_id
        FROM curso_programado cp
        JOIN programa_academico pa ON cp.programa_id = pa.id
        JOIN facultad f ON pa.facultad_id = f.id
        LEFT JOIN docente d ON cp.docente_id = d.id
        WHERE {where_clause}
        ORDER BY f.nombre, pa.tipo, pa.nombre_corto, cp.ciclo, cp.asignatura
    """
    cur.execute(sql_cursos, params)
    cursos = cur.fetchall()

    # =========================
    # Resumen de totales
    # =========================
    sql_total = f"""
        SELECT COALESCE(SUM(cp.remuneracion_monto), 0) AS total_rem
        FROM curso_programado cp
        JOIN programa_academico pa ON cp.programa_id = pa.id
        WHERE {where_clause}
    """
    cur.execute(sql_total, params)
    total_row = cur.fetchone()
    total_remuneracion = total_row["total_rem"] if total_row else 0.0

    total_cursos = len(cursos)

    conn.close()

    return render_template(
        "cursos_programados.html",
        # filtros
        periodo_id=periodo_id,
        facultad_id=facultad_id,
        tipo_programa=tipo_programa,
        tipo_docente_mes=tipo_docente_mes,
        # catálogos
        periodos_global=periodos_global,
        facultades_global=facultades_global,
        tipos_programa_global=tipos_programa_global,
        tipos_docente_mes_global=tipos_docente_mes_global,
        # datos
        periodo_info=periodo_info,
        cursos=cursos,
        total_cursos=total_cursos,
        total_remuneracion=total_remuneracion,
    )


@app.route("/programacion/nuevo", methods=["GET", "POST"])
def programacion_nueva():
    conn = get_db_connection()

    periodos = conn.execute(
        "SELECT id, etiqueta FROM periodo ORDER BY anio DESC, mes DESC"
    ).fetchall()

    facultades = conn.execute(
        "SELECT id, nombre FROM facultad WHERE activo = 1 ORDER BY nombre"
    ).fetchall()

    programas = conn.execute(
        """
        SELECT pa.id, pa.nombre_corto, pa.tipo, f.nombre AS facultad_nombre
        FROM programa_academico pa
        JOIN facultad f ON f.id = pa.facultad_id
        WHERE pa.activo = 1
        ORDER BY f.nombre, pa.tipo DESC, pa.nombre_corto
        """
    ).fetchall()

    docentes = conn.execute(
        "SELECT id, nombre_completo, dni FROM docente WHERE activo = 1 ORDER BY nombre_completo"
    ).fetchall()

    conn.close()

    selected_periodo_id = request.args.get("periodo_id")

    if request.method == "POST":
        periodo_id = request.form.get("periodo_id")
        programa_id = request.form.get("programa_id")
        docente_id_raw = request.form.get("docente_id")  # puede venir "" si sin docente
        ciclo = request.form.get("ciclo", "").strip()
        asignatura = request.form.get("asignatura", "").strip()
        fechas_texto = request.form.get("fechas_texto", "").strip()
        remuneracion_monto_str = request.form.get("remuneracion_monto", "").strip()
        remuneracion_texto = request.form.get("remuneracion_texto", "").strip()
        poi = request.form.get("poi", "").strip()
        tipo_docente_mes = request.form.get("tipo_docente_mes", "").strip() or None
        estado_programacion = (
            request.form.get("estado_programacion", "").strip() or "PROPUESTO"
        )
        observaciones = request.form.get("observaciones", "").strip() or None

        # NUEVOS CAMPOS PARA CARTA
        codigo = request.form.get("codigo", "").strip() or None
        categoria = request.form.get("categoria", "").strip() or None
        sem1 = request.form.get("sem1", "").strip() or None
        sem2 = request.form.get("sem2", "").strip() or None

        # docente_id puede ser None
        docente_id = int(docente_id_raw) if docente_id_raw else None

        # Validación: ahora docente ya no es obligatorio
        if not periodo_id or not programa_id or not asignatura or not fechas_texto:
            flash(
                "Período, programa, asignatura y fechas son obligatorios",
                "danger",
            )
            return render_template(
                "curso_form.html",
                periodos=periodos,
                facultades=facultades,
                programas=programas,
                docentes=docentes,
                curso=None,
                selected_periodo_id=selected_periodo_id,
            )

        try:
            remuneracion_monto = (
                float(remuneracion_monto_str) if remuneracion_monto_str else None
            )
        except ValueError:
            flash("La remuneración debe ser un número válido", "danger")
            return render_template(
                "curso_form.html",
                periodos=periodos,
                facultades=facultades,
                programas=programas,
                docentes=docentes,
                curso=None,
                selected_periodo_id=selected_periodo_id,
            )

        if not remuneracion_texto and remuneracion_monto is not None:
            remuneracion_texto = (
                f"S/. {remuneracion_monto:,.2f}"
                .replace(",", "X")
                .replace(".", ",")
                .replace("X", ".")
            )

        conn = get_db_connection()

        # dni_docente solo si hay docente_id
        dni_docente = None
        if docente_id is not None:
            dni_row = conn.execute(
                "SELECT dni FROM docente WHERE id = ?", (docente_id,)
            ).fetchone()
            dni_docente = dni_row["dni"] if dni_row else None

        conn.execute(
            """
            INSERT INTO curso_programado (
                periodo_id, programa_id, docente_id,
                ciclo, asignatura,
                fechas_texto,
                remuneracion_monto, remuneracion_texto,
                poi, dni_docente,
                tipo_docente_mes, estado_programacion, observaciones,
                codigo, categoria, sem1, sem2
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                periodo_id,
                programa_id,
                docente_id,
                ciclo,
                asignatura,
                fechas_texto,
                remuneracion_monto,
                remuneracion_texto,
                poi,
                dni_docente,
                tipo_docente_mes,
                estado_programacion,
                observaciones,
                codigo,
                categoria,
                sem1,
                sem2,
            ),
        )
        conn.commit()
        conn.close()

        flash("Curso programado creado correctamente ✅", "success")
        return redirect(url_for("programacion", periodo_id=periodo_id))

    return render_template(
        "curso_form.html",
        periodos=periodos,
        facultades=facultades,
        programas=programas,
        docentes=docentes,
        curso=None,
        selected_periodo_id=selected_periodo_id,
    )


@app.route("/programacion/<int:curso_id>/editar", methods=["GET", "POST"])
def programacion_editar(curso_id):
    conn = get_db_connection()

    curso = conn.execute(
        """
        SELECT cp.*, d.nombre_completo AS docente_nombre
        FROM curso_programado cp
        LEFT JOIN docente d ON d.id = cp.docente_id
        WHERE cp.id = ?
        """,
        (curso_id,),
    ).fetchone()

    if not curso:
        conn.close()
        flash("Curso no encontrado", "danger")
        return redirect(url_for("programacion"))

    periodos = conn.execute(
        "SELECT id, etiqueta FROM periodo ORDER BY anio DESC, mes DESC"
    ).fetchall()

    facultades = conn.execute(
        "SELECT id, nombre FROM facultad WHERE activo = 1 ORDER BY nombre"
    ).fetchall()

    programas = conn.execute(
        """
        SELECT pa.id, pa.nombre_corto, pa.tipo, f.nombre AS facultad_nombre
        FROM programa_academico pa
        JOIN facultad f ON f.id = pa.facultad_id
        WHERE pa.activo = 1
        ORDER BY f.nombre, pa.tipo DESC, pa.nombre_corto
        """
    ).fetchall()

    docentes = conn.execute(
        "SELECT id, nombre_completo, dni FROM docente WHERE activo = 1 ORDER BY nombre_completo"
    ).fetchall()

    if request.method == "POST":
        periodo_id = request.form.get("periodo_id")
        programa_id = request.form.get("programa_id")
        docente_id_raw = request.form.get("docente_id")
        ciclo = request.form.get("ciclo", "").strip()
        asignatura = request.form.get("asignatura", "").strip()
        fechas_texto = request.form.get("fechas_texto", "").strip()
        remuneracion_monto_str = request.form.get("remuneracion_monto", "").strip()
        remuneracion_texto = request.form.get("remuneracion_texto", "").strip()
        poi = request.form.get("poi", "").strip()
        tipo_docente_mes = request.form.get("tipo_docente_mes", "").strip() or None
        estado_programacion = (
            request.form.get("estado_programacion", "").strip() or "PROPUESTO"
        )
        observaciones = request.form.get("observaciones", "").strip() or None

        # NUEVOS CAMPOS PARA CARTA
        codigo = request.form.get("codigo", "").strip() or None
        categoria = request.form.get("categoria", "").strip() or None
        sem1 = request.form.get("sem1", "").strip() or None
        sem2 = request.form.get("sem2", "").strip() or None

        docente_id = int(docente_id_raw) if docente_id_raw else None

        if not periodo_id or not programa_id or not asignatura or not fechas_texto:
            flash(
                "Período, programa, asignatura y fechas son obligatorios",
                "danger",
            )
            return render_template(
                "curso_form.html",
                periodos=periodos,
                facultades=facultades,
                programas=programas,
                docentes=docentes,
                curso=curso,
                selected_periodo_id=None,
            )

        try:
            remuneracion_monto = (
                float(remuneracion_monto_str) if remuneracion_monto_str else None
            )
        except ValueError:
            flash("La remuneración debe ser un número válido", "danger")
            return render_template(
                "curso_form.html",
                periodos=periodos,
                facultades=facultades,
                programas=programas,
                docentes=docentes,
                curso=curso,
                selected_periodo_id=None,
            )

        if not remuneracion_texto and remuneracion_monto is not None:
            remuneracion_texto = (
                f"S/. {remuneracion_monto:,.2f}"
                .replace(",", "X")
                .replace(".", ",")
                .replace("X", ".")
            )

        # dni_docente solo si hay docente
        dni_docente = None
        if docente_id is not None:
            dni_row = conn.execute(
                "SELECT dni FROM docente WHERE id = ?", (docente_id,)
            ).fetchone()
            dni_docente = dni_row["dni"] if dni_row else None

        conn.execute(
            """
            UPDATE curso_programado
            SET periodo_id = ?,
                programa_id = ?,
                docente_id = ?,
                ciclo = ?,
                asignatura = ?,
                fechas_texto = ?,
                remuneracion_monto = ?,
                remuneracion_texto = ?,
                poi = ?,
                dni_docente = ?,
                tipo_docente_mes = ?,
                estado_programacion = ?,
                observaciones = ?,
                codigo = ?,
                categoria = ?,
                sem1 = ?,
                sem2 = ?,
                updated_at = datetime('now')
            WHERE id = ?
            """,
            (
                periodo_id,
                programa_id,
                docente_id,
                ciclo,
                asignatura,
                fechas_texto,
                remuneracion_monto,
                remuneracion_texto,
                poi,
                dni_docente,
                tipo_docente_mes,
                estado_programacion,
                observaciones,
                codigo,
                categoria,
                sem1,
                sem2,
                curso_id,
            ),
        )
        conn.commit()
        conn.close()

        flash("Curso programado actualizado correctamente ✅", "success")
        return redirect(url_for("programacion", periodo_id=periodo_id))

    conn.close()
    return render_template(
        "curso_form.html",
        periodos=periodos,
        facultades=facultades,
        programas=programas,
        docentes=docentes,
        curso=curso,
        selected_periodo_id=None,
    )


@app.route("/programacion/<int:curso_id>/eliminar", methods=["POST"])
def programacion_eliminar(curso_id):
    periodo_id = request.form.get("periodo_id")
    facultad_id = request.form.get("facultad_id")
    tipo_programa = request.form.get("tipo_programa", "")
    tipo_docente_mes = request.form.get("tipo_docente_mes", "")

    conn = get_db_connection()
    conn.execute("DELETE FROM curso_programado WHERE id = ?", (curso_id,))
    conn.commit()
    conn.close()

    flash("Curso programado eliminado correctamente ✅", "success")
    return redirect(
        url_for(
            "programacion",
            periodo_id=periodo_id,
            facultad_id=facultad_id,
            tipo_programa=tipo_programa,
            tipo_docente_mes=tipo_docente_mes,
        )
    )


# ===================================================================
# 4) EXPORTACIÓN CSV – ahora incl. código, categoría, sem1, sem2
# ===================================================================

@app.route("/export/programacion.csv")
def export_programacion_csv():
    periodo_id = request.args.get("periodo_id")
    facultad_id = request.args.get("facultad_id")
    tipo_programa = request.args.get("tipo_programa", "")
    tipo_docente_mes = request.args.get("tipo_docente_mes", "")

    conn = get_db_connection()

    if not periodo_id:
        periodo = conn.execute(
            "SELECT id FROM periodo ORDER BY anio DESC, mes DESC LIMIT 1"
        ).fetchone()
        if periodo:
            periodo_id = periodo["id"]
        else:
            periodo_id = None

    params = []
    where_clauses = []

    if periodo_id:
        where_clauses.append("cp.periodo_id = ?")
        params.append(periodo_id)

    if facultad_id:
        where_clauses.append("f.id = ?")
        params.append(facultad_id)

    if tipo_programa:
        where_clauses.append("pa.tipo = ?")
        params.append(tipo_programa)

    if tipo_docente_mes:
        where_clauses.append("cp.tipo_docente_mes = ?")
        params.append(tipo_docente_mes)

    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)

    sql = f"""
        SELECT
            f.nombre AS facultad,
            pa.tipo AS tipo_programa,
            pa.nombre_corto AS programa,
            d.nombre_completo AS docente,
            d.dni AS dni,
            d.tipo_docente AS tipo_docente_global,
            cp.tipo_docente_mes,
            cp.ciclo,
            cp.asignatura,
            cp.fechas_texto,
            cp.remuneracion_texto,
            cp.remuneracion_monto,
            cp.poi,
            cp.codigo,
            cp.categoria,
            cp.sem1,
            cp.sem2,
            p.etiqueta AS periodo
        FROM curso_programado cp
        LEFT JOIN docente d ON d.id = cp.docente_id
        JOIN programa_academico pa ON pa.id = cp.programa_id
        JOIN facultad f ON f.id = pa.facultad_id
        JOIN periodo p ON p.id = cp.periodo_id
        {where_sql}
        ORDER BY
            f.nombre,
            pa.tipo DESC,
            pa.nombre_corto,
            cp.ciclo,
            d.nombre_completo
    """

    rows = conn.execute(sql, params).fetchall()
    conn.close()

    si = StringIO()
    writer = csv.writer(si)
    writer.writerow(
        [
            "Facultad",
            "Tipo programa",
            "Programa",
            "Docente",
            "DNI",
            "Tipo docente (global)",
            "Tipo docente (mes)",
            "Ciclo",
            "Asignatura",
            "Fechas",
            "Remuneración texto",
            "Remuneración monto",
            "POI",
            "Código",
            "Categoría",
            "Sem1",
            "Sem2",
            "Periodo",
        ]
    )
    for r in rows:
        writer.writerow(
            [
                r["facultad"] or "",
                r["tipo_programa"] or "",
                r["programa"] or "",
                r["docente"] or "",
                r["dni"] or "",
                r["tipo_docente_global"] or "",
                r["tipo_docente_mes"] or "",
                r["ciclo"] or "",
                r["asignatura"] or "",
                r["fechas_texto"] or "",
                r["remuneracion_texto"] or "",
                r["remuneracion_monto"] or "",
                r["poi"] or "",
                r["codigo"] or "",
                r["categoria"] or "",
                r["sem1"] or "",
                r["sem2"] or "",
                r["periodo"] or "",
            ]
        )

    output = si.getvalue()
    si.close()

    filename = "programacion_epg.csv"
    return Response(
        output,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ===================================================================
# 5) EXPORTACIÓN A WORD (DOCX) – Formato EPG (listado docentes)
# ===================================================================

def obtener_programa_base(nombre_programa):
    """
    Devuelve el nombre base del programa, sin '2da/3ra... Promoción ...'.
    """
    if not nombre_programa:
        return ""

    base = re.sub(
        r"\s+\d+(?:ra|da|ta|ma|va)\s+Promoción.*$",
        "",
        nombre_programa,
        flags=re.IGNORECASE,
    ).strip(" -–—")

    return base or nombre_programa


FACULTY_ORDER = [
    "Ciencias y Humanidades",
    "Ciencias de la Salud",
    "Ingenierías y Arquitectura",
    "Ciencias Económicas y Contables",
    "Derecho y Ciencia Política",
]


def _nombre_simple_facultad(nombre: str) -> str:
    if not nombre:
        return ""
    nombre = re.sub(r"(?i)^facultad\s+de\s+", "", nombre).strip()
    return nombre


def _facultad_sort_key(nombre: str):
    simple = _nombre_simple_facultad(nombre)
    try:
        idx = FACULTY_ORDER.index(simple)
    except ValueError:
        idx = len(FACULTY_ORDER)
    return (idx, simple.lower())


def _programa_sort_key(programa_base: str):
    if not programa_base:
        return (3, "")
    low = programa_base.lower()
    if low.startswith("doctorado"):
        return (0, low)
    if low.startswith("maestr"):
        return (1, low)
    return (2, low)


def normalizar_programa_base(nombre_programa, tipo_programa):
    if not nombre_programa:
        return ""

    base = nombre_programa

    base = re.sub(
        r"\s+\d+\S*\s+Promoción.*$",
        "",
        base,
        flags=re.IGNORECASE,
    ).strip()

    if re.match(r"(?i)^maestr[ií]a\s+en\s+derecho\b", base):
        return "Maestría en Derecho"

    if re.match(r"(?i)^maestr[ií]a\s+en\s+ingenier[ií]a\s+civil\b", base):
        return "Maestría en Ingeniería Civil"

    return base


def agrupar_por_facultad_y_programa_base(rows):
    grouped = {}
    for r in rows:
        facultad = r["facultad"]
        tipo_prog = (r["tipo_programa"] or "").upper()
        programa = r["programa"] or ""

        programa_base = normalizar_programa_base(programa, tipo_prog)

        if facultad not in grouped:
            grouped[facultad] = {}
        if programa_base not in grouped[facultad]:
            grouped[facultad][programa_base] = []

        grouped[facultad][programa_base].append(r)

    return grouped


@app.route("/export/programacion.docx")
def export_programacion_docx():
    periodo_id = request.args.get("periodo_id")
    facultad_id = request.args.get("facultad_id")
    tipo_programa = request.args.get("tipo_programa", "")
    tipo_docente_mes = request.args.get("tipo_docente_mes", "")

    col_observaciones = request.args.get("col_observaciones") == "1"
    col_universidad = request.args.get("col_universidad") == "1"
    col_telefono = request.args.get("col_telefono") == "1"
    col_correo = request.args.get("col_correo") == "1"
    col_direccion = request.args.get("col_direccion") == "1"

    conn = get_db_connection()

    if not periodo_id:
        periodo_row = conn.execute(
            "SELECT id, etiqueta FROM periodo ORDER BY anio DESC, mes DESC LIMIT 1"
        ).fetchone()
        if periodo_row:
            periodo_id = periodo_row["id"]
        else:
            periodo_id = None

    params = []
    where_clauses = []

    if periodo_id:
        where_clauses.append("cp.periodo_id = ?")
        params.append(periodo_id)

    if facultad_id:
        where_clauses.append("f.id = ?")
        params.append(facultad_id)

    if tipo_programa:
        where_clauses.append("pa.tipo = ?")
        params.append(tipo_programa)

    if tipo_docente_mes:
        where_clauses.append("cp.tipo_docente_mes = ?")
        params.append(tipo_docente_mes)

    where_sql = ""
    if where_clauses:
        where_sql = "WHERE " + " AND ".join(where_clauses)

    sql = f"""
        SELECT
            cp.id AS curso_id,
            f.nombre AS facultad,
            pa.tipo AS tipo_programa,
            pa.nombre_corto AS programa,
            pa.modalidad,
            pa.universidad_procedencia,
            d.nombre_completo AS docente,
            d.dni,
            d.telefono,
            d.correo,
            d.direccion,
            d.tipo_docente AS tipo_docente_global,
            d.titulo_universidad,
            d.magister_universidad,
            d.doctor_universidad,
            cp.tipo_docente_mes,
            cp.ciclo,
            cp.asignatura,
            cp.fechas_texto,
            cp.remuneracion_texto,
            cp.remuneracion_monto,
            cp.poi,
            cp.observaciones,
            p.etiqueta AS periodo_etiqueta
        FROM curso_programado cp
        LEFT JOIN docente d ON d.id = cp.docente_id
        JOIN programa_academico pa ON pa.id = cp.programa_id
        JOIN facultad f ON f.id = pa.facultad_id
        JOIN periodo p ON p.id = cp.periodo_id
        {where_sql}
        ORDER BY
            f.nombre,
            pa.tipo DESC,
            pa.nombre_corto,
            cp.tipo_docente_mes DESC,
            d.nombre_completo,
            cp.ciclo
    """

    rows = conn.execute(sql, params).fetchall()

    periodo_etiqueta = ""
    if rows:
        periodo_etiqueta = rows[0]["periodo_etiqueta"]
    else:
        if periodo_id:
            row = conn.execute(
                "SELECT etiqueta FROM periodo WHERE id = ?", (periodo_id,)
            ).fetchone()
            if row:
                periodo_etiqueta = row["etiqueta"]

    conn.close()

    locales = []
    ordin = []

    for r in rows:
        tipo_mes = (r["tipo_docente_mes"] or "").upper()
        if tipo_mes == "ORDINARIZADO":
            ordin.append(r)
        else:
            locales.append(r)

    def pick_universidad(r):
        for k in ("doctor_universidad", "magister_universidad", "titulo_universidad"):
            val = r[k]
            if val:
                return val
        return r["universidad_procedencia"] or ""

    locales_grouped = agrupar_por_facultad_y_programa_base(locales)

    def _sort_key_ordin(r):
        facultad = r["facultad"] or ""
        programa = r["programa"] or ""
        tipo_prog = (r["tipo_programa"] or "").upper()
        programa_base = normalizar_programa_base(programa, tipo_prog)
        return (
            _facultad_sort_key(facultad),
            _programa_sort_key(programa_base),
            (r["docente"] or "").lower(),
            r["ciclo"] or "",
        )

    ordin_sorted = sorted(ordin, key=_sort_key_ordin)

    doc = Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    titulo = "DOCENTES EPG"
    if periodo_etiqueta:
        titulo += f" – MES {periodo_etiqueta.upper()}"
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(titulo)
    r_title.bold = True
    p_title.alignment = 1

    doc.add_paragraph("")

    def agregar_seccion(doc, encabezado, grouped, start_index):
        n = start_index

        if not grouped:
            return n

        p = doc.add_paragraph()
        r = p.add_run(encabezado)
        r.bold = True
        doc.add_paragraph("")

        facultades_ordenadas = sorted(grouped.keys(), key=_facultad_sort_key)

        for facultad in facultades_ordenadas:
            programas_base = grouped[facultad]

            p_fac = doc.add_paragraph()
            texto_facultad = facultad
            if not facultad.lower().startswith("facultad"):
                texto_facultad = f"Facultad de {facultad}"

            r_fac = p_fac.add_run(texto_facultad)
            r_fac.bold = True

            programas_ordenados = sorted(programas_base.keys(), key=_programa_sort_key)

            for programa_base in programas_ordenados:
                cursos_prog = programas_base[programa_base]

                p_prog = doc.add_paragraph(style="List Bullet")
                r_prog = p_prog.add_run(programa_base)
                r_prog.bold = True

                headers = [
                    "N°",
                    "Docente",
                    "Programa",
                    "Ciclo",
                    "Asignatura",
                    "Fechas de dictado",
                    "Remuneración",
                    "POI",
                    "DNI",
                ]
                extra_cols = []
                if col_observaciones:
                    extra_cols.append("Observaciones")
                if col_universidad:
                    extra_cols.append("Universidad de último grado")
                if col_telefono:
                    extra_cols.append("Teléfono")
                if col_correo:
                    extra_cols.append("Correo")
                if col_direccion:
                    extra_cols.append("Dirección")

                headers.extend(extra_cols)

                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                table.autofit = True

                hdr_cells = table.rows[0].cells
                for i, h in enumerate(headers):
                    hdr_cells[i].text = h

                cursos_prog_ordenados = sorted(
                    cursos_prog,
                    key=lambda r: ((r["ciclo"] or ""), (r["docente"] or "").lower()),
                )

                for rdata in cursos_prog_ordenados:
                    row_cells = table.add_row().cells
                    row_cells[0].text = f"{n:02d}"
                    row_cells[1].text = rdata["docente"] or ""
                    row_cells[2].text = rdata["programa"] or ""
                    row_cells[3].text = rdata["ciclo"] or ""
                    row_cells[4].text = rdata["asignatura"] or ""
                    row_cells[5].text = rdata["fechas_texto"] or ""
                    row_cells[6].text = rdata["remuneracion_texto"] or ""
                    row_cells[7].text = rdata["poi"] or ""
                    row_cells[8].text = rdata["dni"] or ""

                    col_idx = 9
                    if col_observaciones:
                        row_cells[col_idx].text = rdata["observaciones"] or ""
                        col_idx += 1
                    if col_universidad:
                        row_cells[col_idx].text = pick_universidad(rdata)
                        col_idx += 1
                    if col_telefono:
                        row_cells[col_idx].text = rdata["telefono"] or ""
                        col_idx += 1
                    if col_correo:
                        row_cells[col_idx].text = rdata["correo"] or ""
                        col_idx += 1
                    if col_direccion:
                        row_cells[col_idx].text = rdata["direccion"] or ""

                    n += 1

                doc.add_paragraph("")

        return n

    def agregar_seccion_ordinarios(doc, encabezado, rows_ordin, start_index):
        n = start_index

        if not rows_ordin:
            return n

        p = doc.add_paragraph()
        r = p.add_run(encabezado)
        r.bold = True
        doc.add_paragraph("")

        headers = [
            "N°",
            "DOCENTE",
            "PROGRAMA",
            "CICLO",
            "ASIGNATURA",
            "FECHAS",
            "REMUNERACION",
            "POI",
            "DNI",
        ]

        extra_cols = []
        if col_observaciones:
            extra_cols.append("OBSERVACIONES")
        if col_universidad:
            extra_cols.append("Universidad")
        if col_telefono:
            extra_cols.append("Teléfono")
        if col_correo:
            extra_cols.append("Correo")
        if col_direccion:
            extra_cols.append("Dirección")

        headers.extend(extra_cols)

        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = True

        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h

        for rdata in rows_ordin:
            row_cells = table.add_row().cells
            row_cells[0].text = f"{n:02d}"
            row_cells[1].text = rdata["docente"] or ""
            row_cells[2].text = rdata["programa"] or ""
            row_cells[3].text = rdata["ciclo"] or ""
            row_cells[4].text = rdata["asignatura"] or ""
            row_cells[5].text = rdata["fechas_texto"] or ""
            row_cells[6].text = rdata["remuneracion_texto"] or ""
            row_cells[7].text = rdata["poi"] or ""
            row_cells[8].text = rdata["dni"] or ""

            col_idx = 9
            if col_observaciones:
                row_cells[col_idx].text = rdata["observaciones"] or ""
                col_idx += 1
            if col_universidad:
                row_cells[col_idx].text = pick_universidad(rdata)
                col_idx += 1
            if col_telefono:
                row_cells[col_idx].text = rdata["telefono"] or ""
                col_idx += 1
            if col_correo:
                row_cells[col_idx].text = rdata["correo"] or ""
                col_idx += 1
            if col_direccion:
                row_cells[col_idx].text = rdata["direccion"] or ""

            n += 1

        doc.add_paragraph("")
        return n

    contador = 1
    contador = agregar_seccion(doc, "DOCENTES LOCALES", locales_grouped, contador)
    if ordin_sorted:
        contador = agregar_seccion_ordinarios(
            doc,
            "DOCENTES ORDINARIZADOS",
            ordin_sorted,
            contador,
        )

    if not rows:
        doc.add_paragraph(
            "No hay cursos programados para el período seleccionado."
        )

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)

    filename = "Docentes_EPG"
    if periodo_etiqueta:
        filename += "_" + periodo_etiqueta.replace(" ", "_")
    filename += ".docx"

    return send_file(
        bio,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


# ===================================================================
# 6) CARTAS DE INVITACIÓN A DOCENTES (desde BD)
# ===================================================================

@app.route("/cartas_invitacion", methods=["GET", "POST"])
def cartas_invitacion():
    """
    Vista para:
    - Elegir período
    - Ver cursos programados y estado de datos para carta
    - Seleccionar cursos y generar cartas .docx/.pdf (ZIP) usando la plantilla Word
    """
    conn = get_db_connection()
    cur = conn.cursor()

    # Períodos para el combo
    periodos = cur.execute(
        "SELECT id, etiqueta FROM periodo ORDER BY anio DESC, mes DESC"
    ).fetchall()

    # Determinar período seleccionado
    periodo_id = request.args.get("periodo_id") or request.form.get("periodo_id")

    if request.method == "GET":
        cursos = []
        if periodo_id:
            cursos = cur.execute(
                """
                SELECT cp.*,
                       d.nombre_completo AS docente_nombre,
                       pa.nombre_corto   AS programa_nombre,
                       pa.tipo           AS programa_tipo,
                       pa.modalidad      AS programa_modalidad
                FROM curso_programado cp
                JOIN programa_academico pa ON pa.id = cp.programa_id
                LEFT JOIN docente d ON d.id = cp.docente_id
                WHERE cp.periodo_id = ?
                ORDER BY programa_nombre, asignatura
                """,
                (periodo_id,),
            ).fetchall()

        conn.close()
        return render_template(
            "cartas_invitacion.html",
            periodos=periodos,
            periodo_id=periodo_id,
            cursos=cursos,
        )

    # POST: generar cartas
    if not PLANTILLA_CARTA.exists():
        conn.close()
        flash(
            f"No se encuentra la plantilla de carta: {PLANTILLA_CARTA.name}. "
            f"Colócala en la misma carpeta que app.py.",
            "danger",
        )
        return redirect(url_for("cartas_invitacion", periodo_id=periodo_id))

    selected_ids = request.form.getlist("curso_id")
    if not selected_ids:
        conn.close()
        flash("No seleccionaste ningún curso para generar cartas.", "warning")
        return redirect(url_for("cartas_invitacion", periodo_id=periodo_id))

    # N° de carta inicial (opcional)
    start_num_str = request.form.get("start_num") or ""
    try:
        next_num = int(start_num_str) if start_num_str else None
    except ValueError:
        next_num = None

    # Info del período (para carpeta y MES)
    periodo_row = None
    if periodo_id:
        periodo_row = cur.execute(
            "SELECT id, etiqueta, anio, mes FROM periodo WHERE id = ?",
            (periodo_id,),
        ).fetchone()

    # Obtener cursos seleccionados
    placeholders = ",".join("?" for _ in selected_ids)
    params = selected_ids

    rows = cur.execute(
        f"""
        SELECT cp.*,
               d.nombre_completo AS docente_nombre,
               pa.nombre_corto   AS programa_nombre,
               pa.tipo           AS programa_tipo,
               pa.modalidad      AS programa_modalidad,
               pe.anio           AS periodo_anio,
               pe.mes            AS periodo_mes,
               pe.etiqueta       AS periodo_etiqueta
        FROM curso_programado cp
        JOIN programa_academico pa ON pa.id = cp.programa_id
        JOIN periodo pe ON pe.id = cp.periodo_id
        LEFT JOIN docente d ON d.id = cp.docente_id
        WHERE cp.id IN ({placeholders})
        ORDER BY cp.id
        """,
        params,
    ).fetchall()

    if not rows:
        conn.close()
        flash("No se encontraron cursos para generar cartas.", "danger")
        return redirect(url_for("cartas_invitacion", periodo_id=periodo_id))

    # Carpeta específica del período, ej: "Cartas de invitación Noviembre 2025"
    if periodo_row:
        etiqueta_periodo = periodo_row["etiqueta"]
    else:
        etiqueta_periodo = rows[0]["periodo_etiqueta"]

    subdir_name = f"Cartas de invitación {etiqueta_periodo}"
    output_dir = CARTAS_OUTPUT_DIR / subdir_name
    output_dir.mkdir(parents=True, exist_ok=True)

    from io import BytesIO
    import zipfile

    zip_buffer = BytesIO()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for row in rows:
            # ======== Datos base del curso/docente =========
            docente = row["docente_nombre"] or "DOCENTE POR DEFINIR"
            asignatura = row["asignatura"]
            programa_nombre = row["programa_nombre"] or ""
            programa_tipo = (row["programa_tipo"] or "").upper()
            programa_modalidad = (row["programa_modalidad"] or "").strip()

            # Programa base (sin "15va promoción...") y texto final
            programa_base = normalizar_programa_base(programa_nombre, programa_tipo)
            if programa_tipo == "DOCTORADO":
                programa_texto = f"del {programa_base}"
            elif programa_tipo == "MAESTRIA":
                programa_texto = f"de la {programa_base}"
            else:
                programa_texto = programa_base

            # Modalidad
            mod_up = programa_modalidad.upper()
            if mod_up == "DISTANCIA":
                modalidad_texto = "MODALIDAD A DISTANCIA"
            elif mod_up:
                modalidad_texto = f"MODALIDAD {mod_up}"
            else:
                modalidad_texto = ""

            # Campos nuevos para carta: codigo / categoria / sem1 / sem2
            codigo = row["codigo"] or ""
            categoria = row["categoria"] or ""
            sem1 = row["sem1"] or ""
            sem2 = row["sem2"] or ""

            # Fecha larga y año ("23 de octubre de 2025")
            fecha_larga, anio = fecha_larga_es(None)
            ciudad = "Cusco"

            # Mes en la carta (ej. "NOVIEMBRE 2025")
            periodo_etiqueta_row = row["periodo_etiqueta"] or etiqueta_periodo
            mes_tabla = (periodo_etiqueta_row or "").upper()

            # Remuneración (numérica -> formato 5,900.00)
            override = None
            if row["remuneracion_monto"] is not None:
                override = str(row["remuneracion_monto"])
            remuneracion_num = calc_remuneracion(programa_texto, override)
            remuneracion = remuneracion_num

            # Número de carta
            if next_num is not None:
                numero = str(next_num)
                next_num += 1
            else:
                numero = "000"

            titulo = ""

            # ======== Contexto para la plantilla DOCX (docxtpl) =========
            context = {
                "ciudad": ciudad,
                "fecha_larga": fecha_larga,
                "numero": numero,
                "anio": anio,
                "titulo": titulo,
                "docente": docente,
                "asignatura": asignatura,
                "programa": programa_texto,
                "modalidad": modalidad_texto,
                "codigo": codigo,
                "categoria": categoria,
                "cred": "04",
                "mes": mes_tabla,
                "sem1": sem1,
                "sem2": sem2,
                "horarios": "VIERNES: 5-10 PM; SABADO: 8-1 PM Y 4 -9 PM; DOMINGO: 8-1 PM",
                "remuneracion": remuneracion,
            }

            # Renderizar plantilla
            tpl = DocxTemplate(str(PLANTILLA_CARTA))
            tpl.render(context)

            # Nombre base de archivo “bonito”
            paterno, nombre = extraer_paterno_y_nombre(docente)
            file_stub = f"CARTA N°{numero} {paterno} {nombre}"
            filename_docx = limpiar_nombre_archivo(file_stub) + ".docx"

            # Guardar DOCX en disco
            out_path = output_dir / filename_docx
            tpl.save(out_path)

            # Añadir DOCX al ZIP
            with open(out_path, "rb") as f:
                zf.writestr(filename_docx, f.read())

            # ======== Intentar generar PDF también y añadirlo al ZIP ========
            if TRY_PDF_WEB:
                try:
                    out_pdf_path = out_path.with_suffix(".pdf")
                    docx2pdf_convert(str(out_path), str(out_pdf_path))

                    with open(out_pdf_path, "rb") as fpdf:
                        zf.writestr(out_pdf_path.name, fpdf.read())
                except Exception as e:
                    # No romper el flujo si falla PDF; solo loggea en consola
                    print(f"[Aviso] No se pudo generar PDF para {out_path.name}: {e}")

    conn.close()

    zip_buffer.seek(0)
    zip_filename = f"cartas_invitacion_{etiqueta_periodo.replace(' ', '_')}.zip"

    flash(
        f"Cartas generadas correctamente en la carpeta '{subdir_name}' "
        f"y descargadas como ZIP (DOCX{', PDF' if TRY_PDF_WEB else ''}) ✅",
        "success",
    )
    return send_file(
        zip_buffer,
        as_attachment=True,
        download_name=zip_filename,
        mimetype="application/zip",
    )


if __name__ == "__main__":
    app.run(debug=True)
